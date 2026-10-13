"""Deterministic T-c check (SPEC §15.1, T2): does an agent-made xlsx list exactly the planted
mismatches with the right numbers? Layout-agnostic: a row belongs to a (warehouse, model) pair
when exactly one warehouse name and one model name appear in its text cells.
    uv run python -m evals.checks.tc <xlsx> <truth.json>"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from fixtures.generate import MODELS, WAREHOUSES
from usta.lang import az_casefold


def _rows(path: Path) -> list[tuple[str, set[float]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            text = " | ".join(az_casefold(v) for v in row if isinstance(v, str))
            nums = {float(v) for v in row if isinstance(v, int | float) and not isinstance(v, bool)}
            rows.append((text, nums))
    wb.close()
    return rows


def check_tc(
    xlsx: Path,
    truth: list[dict],
    warehouses: list[str] = WAREHOUSES,
    models: list[str] = MODELS,
) -> list[str]:
    found: dict[tuple[str, str], list[set[float]]] = {}
    for text, nums in _rows(xlsx):
        whs = [w for w in warehouses if az_casefold(w) in text]
        ms = [m for m in models if az_casefold(m) in text]
        if len(whs) == 1 and len(ms) == 1:
            found.setdefault((whs[0], ms[0]), []).append(nums)
    expected = {(t["anbar"], t["model"]): t for t in truth}
    problems = []
    for key, t in expected.items():
        if key not in found:
            problems.append(f"missing mismatch: {key[0]} / {key[1]}")
        elif not any({float(t["siyahida"]), float(t["qaliq"])} <= nums for nums in found[key]):
            problems.append(
                f"wrong numbers for {key[0]} / {key[1]}: expected list {t['siyahida']} "
                f"and stock {t['qaliq']}"
            )
    for key in sorted(set(found) - set(expected)):
        problems.append(f"not a mismatch: {key[0]} / {key[1]}")
    return problems


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: python -m evals.checks.tc <xlsx> <truth.json>")
    issues = check_tc(Path(sys.argv[1]), json.loads(Path(sys.argv[2]).read_text(encoding="utf-8")))
    print("\n".join(issues) or "T-c check passed")
    sys.exit(1 if issues else 0)
