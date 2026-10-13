"""Excel writing helper for the Usta sandbox (openpyxl)."""

from __future__ import annotations

import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(name: str) -> str:
    return _BAD_SHEET_CHARS.sub("_", name)[:31] or "Sheet"


def write_xlsx(path: str, sheets: dict[str, list[list]]) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=_sheet_title(name))
        for row in rows:
            ws.append(list(row))
        if rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EEF5")
            ws.freeze_panes = "A2"
        for column in ws.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=0)
            ws.column_dimensions[get_column_letter(column[0].column)].width = min(60, width + 2)
    wb.save(path)
    return path
