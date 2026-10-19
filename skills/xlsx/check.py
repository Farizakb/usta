"""check.py for xlsx (SPEC §7.1, §9): python check.py <file.xlsx>
Prints one JSON line {"ok", "reasons", "text_sample"}; exit 0 when checked, 2 on usage error."""

import json
import sys

ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
SAMPLE_CHARS = 4000


def emit(ok: bool, reasons: list[str], sample: str) -> None:
    print(json.dumps({"ok": ok, "reasons": reasons, "text_sample": sample}, ensure_ascii=False))


def check(path: str) -> None:
    from openpyxl import load_workbook

    reasons: list[str] = []
    texts: list[str] = []
    values = 0
    data_rows = 0
    try:
        wb = load_workbook(path, read_only=True)
        try:
            for ws in wb.worksheets:
                seen_header = False
                for row in ws.iter_rows():
                    cells = [c for c in row if c.value is not None and c.value != ""]
                    if not cells:
                        continue
                    if seen_header:
                        data_rows += 1
                    seen_header = True
                    for c in cells:
                        v = c.value
                        values += 1
                        if isinstance(v, str):
                            if v.strip().upper() in ERRORS:
                                reasons.append(f"sheet {ws.title!r} has an error value {v.strip()}")
                            elif v.startswith("="):
                                # no formula engine in the sandbox: an unevaluated formula
                                # cannot be checked (#DIV/0!, #REF!), so it fails (P8)
                                reasons.append(
                                    f"sheet {ws.title!r} cell {c.coordinate} has a formula; "
                                    "write computed values"
                                )
                            else:
                                texts.append(v)
        finally:
            wb.close()
    except Exception as e:
        emit(False, [f"cannot read workbook: {type(e).__name__}: {e}"], "")
        return
    if values == 0:
        reasons.append("workbook has no data")
    elif data_rows == 0:
        reasons.append("workbook has no data rows (only header rows)")
    text = " | ".join(texts)
    if "�" in text or "\x00" in text:
        reasons.append("text contains U+FFFD or NUL characters (broken encoding)")
    emit(not reasons, sorted(set(reasons)), text[:SAMPLE_CHARS])


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("usage: check.py <file.xlsx>", file=sys.stderr)
        sys.exit(2)
    check(sys.argv[1])
