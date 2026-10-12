import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from evals.checks.tc import check_tc
from fixtures.generate import MODELS, N_MISMATCHES, WAREHOUSES, build, write


def test_build_is_deterministic_and_plants_the_mismatches() -> None:
    devices, stock, truth = build()
    assert build() == (devices, stock, truth)
    assert devices[0] == ["IMEI", "Model", "Anbar", "Status"]
    assert stock[0] == ["Anbar", "Model", "Qalıq"]
    assert len(truth) == N_MISMATCHES
    assert all(t["ferq"] != 0 and t["qaliq"] >= 0 for t in truth)
    assert 15 <= len(devices) - 1 <= 60
    assert len({row[0] for row in devices[1:]}) == len(devices) - 1  # unique IMEIs
    listed = {(w, m): 0 for w in WAREHOUSES for m in MODELS}
    for _, model, wh, _ in devices[1:]:
        listed[(wh, model)] += 1
    for wh, model, qty in stock[1:]:
        planted = [t for t in truth if (t["anbar"], t["model"]) == (wh, model)]
        assert (qty != listed[(wh, model)]) == bool(planted)


def test_write_puts_truth_outside_the_workspace(tmp_path: Path) -> None:
    files, truth_path = tmp_path / "files", tmp_path / "secret" / "truth.json"
    write(files, truth_path)
    assert sorted(p.name for p in files.iterdir()) == ["anbar_qaliqlari.xlsx", "cihazlar.xlsx"]
    assert len(json.loads(truth_path.read_text(encoding="utf-8"))) == N_MISMATCHES
    ws = load_workbook(files / "anbar_qaliqlari.xlsx").active
    assert [c.value for c in ws[1]] == ["Anbar", "Model", "Qalıq"]


def table(tmp_path: Path, rows: list[list]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["Anbar", "Model", "Siyahıda", "Qalıq", "Fərq"])
    for r in rows:
        ws.append(r)
    path = tmp_path / "out.xlsx"
    wb.save(path)
    return path


def rows_for(truth: list[dict]) -> list[list]:
    return [[t["anbar"], t["model"], t["siyahida"], t["qaliq"], t["ferq"]] for t in truth]


def test_correct_table_passes(tmp_path: Path) -> None:
    _, _, truth = build()
    assert check_tc(table(tmp_path, rows_for(truth)), truth) == []


def az_upper(text: str) -> str:
    # str.upper() maps both i and ı to I, which loses the distinction; Azerbaijani does not.
    return text.replace("i", "İ").replace("ı", "I").upper()


def test_case_and_dotted_i_do_not_matter(tmp_path: Path) -> None:
    _, _, truth = build()
    rows = [[az_upper(r[0]), r[1].lower(), *r[2:]] for r in rows_for(truth)]
    assert check_tc(table(tmp_path, rows), truth) == []


def test_missing_extra_and_wrong_numbers_are_reported(tmp_path: Path) -> None:
    _, _, truth = build()
    rows = rows_for(truth)
    missing = rows[0]
    rows = rows[1:]
    rows[0] = [rows[0][0], rows[0][1], 99, 98, 1]
    planted = {(t["anbar"], t["model"]) for t in truth}
    extra = next((w, m) for w in WAREHOUSES for m in MODELS if (w, m) not in planted)
    rows.append([extra[0], extra[1], 2, 2, 0])
    problems = check_tc(table(tmp_path, rows), truth)
    assert f"missing mismatch: {missing[0]} / {missing[1]}" in problems
    assert any(p.startswith("wrong numbers for") for p in problems)
    assert f"not a mismatch: {extra[0]} / {extra[1]}" in problems
